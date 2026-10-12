#!/usr/bin/env python3
"""Tests for review_sheet.py. All words below are TEST STRINGS ONLY (arbitrary code point sequences
or SYNTH_* placeholders); none is Malayalam vocabulary and none is evidence of anything."""
import csv
import io
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
try:
    import openpyxl  # noqa: F401
except ImportError:
    print("review_sheet tests skipped: openpyxl not installed")
    sys.exit(0)
from openpyxl import load_workbook

import review_sheet as rs

BUILD = os.environ.get("S2W_BUILD", "build")
# arbitrary Malayalam-block code point strings (mechanics only)
A, B, C, D = "കക", "കഖ", "കഗ", "കഘ"
CONJ, CONJ_ZWNJ = "ക്ഷ", "ക്‌ഷ"


def run(*argv):
    out, err = io.StringIO(), io.StringIO()
    code = 0
    with redirect_stdout(out), redirect_stderr(err):
        try:
            rs.main(list(argv))
        except SystemExit as e:
            code = e.code if isinstance(e.code, int) else 1
            if isinstance(e.code, str):
                err.write(e.code)
    return code, out.getvalue(), err.getvalue()


def row(**kw):
    d = {k: "" for k in rs.KEYS}
    d.update(id="W001", role="vocab", ctype="none", word=A)
    d.update(kw)
    return d


def full(**kw):
    base = dict(r1_dec="Accept", r2_dec="Accept", r1_fam="Yes", r2_fam="Yes")
    base.update(kw)
    return row(**base)


class Logic(unittest.TestCase):
    def st(self, **kw):
        return rs.evaluate_row(full(**kw))

    def test_basic_states(self):
        self.assertEqual(rs.evaluate_row(row(word=""))[0], "Empty slot")
        self.assertEqual(rs.evaluate_row(row())[0], "Pending review")
        self.assertEqual(rs.evaluate_row(row(r1_dec="Accept"))[0], "Pending review")
        self.assertEqual(self.st(), ("Agreed: accept", A))
        self.assertEqual(self.st(r1_dec="Reject", r2_dec="Reject")[0], "Agreed: reject")
        self.assertTrue(self.st(r1_dec="Reject")[0].startswith("UNRESOLVED"))
        self.assertTrue(rs.evaluate_row(row(r1_dec="Accept", r2_dec="Accept"))[0] == "Pending review")  # familiarity missing

    def test_correction_rules(self):
        self.assertEqual(self.st(r1_dec="Correct", r1_fix=B, r2_dec="Correct", r2_fix=B), ("Agreed: accept", B))
        self.assertTrue(self.st(r1_dec="Correct", r1_fix=B)[0].startswith("UNRESOLVED"))       # accept vs correct
        self.assertTrue(self.st(r1_dec="Correct", r1_fix=B, r2_dec="Correct", r2_fix=C)[0].startswith("UNRESOLVED"))
        self.assertTrue(self.st(r1_dec="Correct")[0].startswith("Invalid entry"))             # no corrected spelling
        self.assertTrue(self.st(r1_dec="Maybe")[0].startswith("Invalid entry"))

    def test_zwj_is_not_ignored(self):
        s = self.st(word=CONJ, r1_dec="Accept", r2_dec="Correct", r2_fix=CONJ_ZWNJ)[0]
        self.assertTrue(s.startswith("UNRESOLVED") and "ZWJ/ZWNJ" in s)
        self.assertEqual(self.st(word=CONJ, r2_dec="Correct", r2_fix=CONJ)[0], "Agreed: accept")

    def test_contrast_rules(self):
        pair = dict(pair="P01", ctype="vowel_length", r1_con="Yes", r2_con="Yes")
        self.assertEqual(self.st(**pair)[0], "Agreed: accept")
        self.assertTrue(self.st(**dict(pair, r2_con="No"))[0].startswith("UNRESOLVED"))
        self.assertTrue(self.st(**dict(pair, r1_con="No", r2_con="No"))[0].startswith("Agreed: contrast not genuine"))
        self.assertEqual(rs.evaluate_row(full(pair="P01", ctype="vowel_length", r1_con="Yes"))[0], "Pending review")
        self.assertTrue(self.st(**dict(pair, r1_con="N/A"))[0].startswith("Invalid entry"))
        self.assertTrue(self.st(pair="P01")[0].startswith("Invalid entry"))                    # pair without type
        self.assertTrue(self.st(ctype="gemination")[0].startswith("Invalid entry"))            # type without pair
        self.assertTrue(self.st(role="oov", **pair)[0].startswith("Invalid entry"))
        self.assertEqual(self.st(r1_con="No")[0], "Agreed: accept")                            # ignored outside a pair

    def test_familiarity_and_ids(self):
        self.assertTrue(self.st(r1_fam="No")[0].startswith("UNRESOLVED"))
        self.assertTrue(self.st(r1_fam="No", r2_fam="No")[0].startswith("Agreed: not familiar"))
        self.assertTrue(self.st(id="bad id")[0].startswith("Invalid entry"))
        self.assertTrue(self.st(role="maybe")[0].startswith("Invalid entry"))
        self.assertTrue(rs.evaluate_row(row(word="", r1_dec="Accept"))[0].startswith("Invalid entry"))

    def test_malayalam_text_check(self):
        self.assertTrue(rs.is_malayalam_text(A) and rs.is_malayalam_text(CONJ_ZWNJ) and rs.is_malayalam_text(A + " " + B))
        self.assertFalse(rs.is_malayalam_text("abc") or rs.is_malayalam_text(A + "x") or rs.is_malayalam_text(""))


class Workbooks(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp()
        self.m = os.path.join(self.d, "master.xlsx")
        self.assertEqual(run("template", self.m)[0], 0)

    def tearDown(self):
        shutil.rmtree(self.d, ignore_errors=True)

    def cell(self, path, sheet, wid, key, value=None, hdr_row=1):
        wb = load_workbook(path)
        ws = wb[sheet]
        idcol = [c.column for c in ws[1] if c.value == rs.HEAD["id"]][0]
        col = [c.column for c in ws[1] if c.value == rs.HEAD[key]][0]
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=idcol).value == wid:
                if value is None:
                    return ws.cell(row=r, column=col).value
                ws.cell(row=r, column=col, value=value)
                wb.save(path)
                return
        raise KeyError(wid)

    def fill(self, path, sheet, wid, **kv):
        wb = load_workbook(path)
        ws = wb[sheet]
        hdr = {c.value: c.column for c in ws[1]}
        idcol = hdr[rs.HEAD["id"]]
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=idcol).value == wid:
                for k, v in kv.items():
                    ws.cell(row=r, column=hdr[rs.HEAD[k]], value=v)
                wb.save(path)
                return
        raise KeyError(wid)

    def test_template_structure(self):
        wb = load_workbook(self.m)
        self.assertEqual(wb.sheetnames, ["README", "Master", "_meta"])
        self.assertEqual(wb["_meta"].sheet_state, "hidden")
        ws = wb["Master"]
        self.assertEqual([c.value for c in ws[1]], [rs.HEAD[k] for k in rs.KEYS])
        self.assertEqual(ws.max_row, 51)
        self.assertEqual((ws["A2"].value, ws["A41"].value, ws["A42"].value, ws["B42"].value), ("W001", "W040", "U001", "oov"))
        for want in ("Word ID", "Malayalam word", "English meaning", "Pronunciation (IPA)", "Contrast pair ID", "Contrast type",
                     "Reviewer 1 decision", "Reviewer 1 notes", "Reviewer 2 decision", "Final verified spelling", "regional variants", "Recording approved"):
            self.assertTrue(any(want.lower() in str(c.value).lower() for c in ws[1]), want)
        self.assertGreaterEqual(len(ws.data_validations.dataValidation), 8)
        self.assertTrue(all(c.value is None for c in ws[2][2:7]))                  # no invented content
        self.assertTrue(all(ws.cell(row=r, column=3).value is None for r in range(2, 52)))
        self.assertEqual(run("template", self.m)[0] != 0, True)                    # never overwrites
        wb2 = self.m + "2.xlsx"
        self.assertEqual(run("template", wb2, "--vocab-slots", "30", "--oov-slots", "0")[0], 0)
        self.assertEqual(load_workbook(wb2)["Master"].max_row, 31)

    def propose(self):
        P = [("W001", A, "P01", "vowel_length"), ("W002", B, "P01", "vowel_length"), ("W003", C, "", "none"), ("W004", D, "", "none"),
             ("W005", CONJ, "", "none"), ("W006", A + B, "", "none"), ("W007", B + C, "P02", "gemination"), ("W008", C + D, "P02", "gemination")]
        for wid, w, pair, ct in P:
            self.fill(self.m, "Master", wid, word=w, meaning="test meaning " + wid, pair=pair or None, ctype=ct)
        self.fill(self.m, "Master", "U001", word=D + D, meaning="test unknown")

    def review(self, path, sheet, n, answers):
        for wid, kv in answers.items():
            self.fill(path, sheet, wid, **{"r%d_%s" % (n, k): v for k, v in kv.items()})

    def test_full_flow(self):
        self.propose()
        out = os.path.join(self.d, "copies")
        code, o, e = run("split", self.m, "--out-dir", out)
        self.assertEqual(code, 0, e)
        c1, c2 = (os.path.join(out, "review_copy_reviewer%d.xlsx" % n) for n in (1, 2))
        # blind copies hold only the proposal + that reviewer's own columns
        for n, path in ((1, c1), (2, c2)):
            ws = load_workbook(path)["Review"]
            heads = [c.value for c in ws[1]]
            self.assertEqual(len(heads), 8 + 6)
            self.assertFalse(any("Reviewer %d" % (3 - n) in str(h) for h in heads))
            self.assertFalse(any("status" in str(h).lower() or "Final" in str(h) for h in heads))
            self.assertTrue(ws.protection.sheet)
            self.assertEqual(ws.max_row, 1 + 9)                                   # only proposed rows
        self.assertNotEqual(run("split", self.m, "--out-dir", out)[0], 0)         # no overwrite

        ok = dict(dec="Accept", fam="Yes", con="Yes")
        r1 = {"W001": ok, "W002": ok, "W003": dict(dec="Correct", fix=B, fam="Yes"), "W004": dict(dec="Correct", fix=C, fam="Yes"),
              "W005": dict(dec="Reject", fam="Yes"), "W006": dict(dec="Reject"), "W007": ok, "W008": ok, "U001": dict(dec="Accept", fam="Yes")}
        r2 = {"W001": ok, "W002": ok, "W003": dict(dec="Accept", fam="Yes"), "W004": dict(dec="Correct", fix=C, fam="Yes"),
              "W005": dict(dec="Accept", fam="Yes"), "W006": dict(dec="Reject"), "W007": dict(dec="Accept", fam="Yes", con="No"), "W008": ok,
              "U001": dict(dec="Accept", fam="Yes")}
        self.review(c1, "Review", 1, r1)
        self.review(c2, "Review", 2, r2)

        # reviewer 2 edits a proposal cell -> merge refuses
        bad = os.path.join(self.d, "tamper.xlsx")
        shutil.copy(c2, bad)
        self.fill(bad, "Review", "W003", word=A)
        code, o, e = run("merge", self.m, "--r1", c1, "--r2", bad, "--out", os.path.join(self.d, "x.xlsx"))
        self.assertNotEqual(code, 0)
        self.assertIn("edited", e)
        # copy swapped between slots -> refused
        self.assertNotEqual(run("merge", self.m, "--r1", c2, "--r2", c1, "--out", os.path.join(self.d, "y.xlsx"))[0], 0)

        merged = os.path.join(self.d, "merged.xlsx")
        code, o, e = run("merge", self.m, "--r1", c1, "--r2", c2, "--out", merged)
        self.assertEqual(code, 0, e)
        stt = {w: self.cell(merged, "Master", w, "status") for w in ("W001", "W003", "W004", "W005", "W006", "W007", "U001")}
        self.assertEqual(stt["W001"], "Agreed: accept")
        self.assertTrue(stt["W003"].startswith("UNRESOLVED"))                      # accept vs correct
        self.assertEqual(stt["W004"], "Agreed: accept")
        self.assertTrue(stt["W005"].startswith("UNRESOLVED"))                      # reject vs accept
        self.assertEqual(stt["W006"], "Agreed: reject")
        self.assertTrue(stt["W007"].startswith("UNRESOLVED"))                      # contrast disagreement
        self.assertEqual(stt["U001"], "Agreed: accept")
        self.assertIn("UNRESOLVED", o)

        # export refuses while unresolved/pending
        wl = os.path.join(self.d, "wordlist.csv")
        code, o, e = run("export", merged, "--out", wl, "--reviewer1-id", "rev-a", "--reviewer2-id", "rev-b")
        self.assertEqual(code, 1)
        self.assertFalse(os.path.exists(wl))
        self.assertIn("UNRESOLVED", e)

        # humans settle: reviewers update the Master answers, record how, and fill the final columns
        self.fill(merged, "Master", "W003", r1_dec="Accept", resolution="settled at meeting")
        self.fill(merged, "Master", "W005", r2_dec="Reject", resolution="settled at meeting")
        self.fill(merged, "Master", "W007", r2_con="Yes", resolution="contrast confirmed by both")
        code, o, e = run("status", merged, "--write")
        self.assertEqual(code, 0)
        self.assertEqual(self.cell(merged, "Master", "W003", "status"), "Agreed: accept")
        # export still refuses: final spellings and approvals are human decisions, not assumed
        code, o, e = run("export", merged, "--out", wl, "--reviewer1-id", "rev-a", "--reviewer2-id", "rev-b")
        self.assertEqual(code, 1)
        self.assertIn("Recording approved", e)
        finals = {"W001": A, "W002": B, "W003": C, "W004": C, "W007": B + C, "W008": C + D, "U001": D + D}
        for wid, sp in finals.items():
            self.fill(merged, "Master", wid, final=sp, approved="Yes", variants="v-" + wid)
        self.fill(merged, "Master", "W005", approved="No")
        self.fill(merged, "Master", "W006", approved="No")
        # W004: both corrected to C; W003 accepted -> candidate is C ... make a wrong final to prove exactness is enforced
        self.fill(merged, "Master", "W004", final=D)
        code, o, e = run("export", merged, "--out", wl, "--reviewer1-id", "rev-a", "--reviewer2-id", "rev-b")
        self.assertEqual(code, 1)
        self.assertIn("must equal the agreed spelling", e)
        self.fill(merged, "Master", "W004", final=C)
        # a pair with a member rejected blocks export
        self.fill(merged, "Master", "W008", approved="No")
        code, o, e = run("export", merged, "--out", wl, "--reviewer1-id", "rev-a", "--reviewer2-id", "rev-b")
        self.assertEqual(code, 1)
        self.assertIn("pair P02", e)
        self.fill(merged, "Master", "W008", approved="Yes")
        self.assertNotEqual(run("export", merged, "--out", wl, "--reviewer1-id", "rev-a", "--reviewer2-id", "rev-a")[0], 0)
        code, o, e = run("export", merged, "--out", wl, "--reviewer1-id", "rev-a", "--reviewer2-id", "rev-b")
        self.assertEqual(code, 0, e)
        with open(wl, encoding="utf-8") as f:
            first = f.readline()
            rows = list(csv.DictReader(f))
        self.assertEqual(first, "#s2w-wordlist v1\n")
        ids = [r["word_id"] for r in rows]
        self.assertEqual(ids, ["W001", "W002", "W003", "W004", "W007", "W008", "U001"])  # rejected W005/W006 absent; vocab first
        self.assertTrue(all(r["review_status"] == "verified" and r["reviewer_1"] == "rev-a" and r["reviewer_2"] == "rev-b" for r in rows))
        self.assertEqual(rows[0]["pair_id"], "P01")
        self.assertEqual(rows[-1]["role"], "oov")
        import json
        with open(os.path.splitext(wl)[0] + ".provenance.json", encoding="utf-8") as pf:
            prov = json.load(pf)
        self.assertEqual((prov["n_included"], prov["n_vocab"], prov["n_oov"]), (7, 6, 1))
        self.assertEqual({e["word_id"] for e in prov["excluded"]}, {"W005", "W006"})
        self.assertEqual(len(prov["master_sha256"]), 64)
        self.assertNotEqual(run("export", merged, "--out", wl, "--reviewer1-id", "rev-a", "--reviewer2-id", "rev-b")[0], 0)   # no overwrite

    def test_proposal_change_after_split_is_detected(self):
        self.propose()
        out = os.path.join(self.d, "copies")
        run("split", self.m, "--out-dir", out)
        self.fill(self.m, "Master", "W003", word=D + A)
        c1, c2 = (os.path.join(out, "review_copy_reviewer%d.xlsx" % n) for n in (1, 2))
        code, o, e = run("merge", self.m, "--r1", c1, "--r2", c2, "--out", os.path.join(self.d, "m.xlsx"))
        self.assertNotEqual(code, 0)
        self.assertIn("different proposal", e)

    def test_non_malayalam_blocked_unless_flagged(self):
        self.fill(self.m, "Master", "W001", word="latinword", meaning="m")
        self.fill(self.m, "Master", "W001", r1_dec="Accept", r1_fam="Yes", r2_dec="Accept", r2_fam="Yes",
                  final="latinword", approved="Yes")
        wl = os.path.join(self.d, "w.csv")
        # only W001 is populated; others are empty slots (ignored); need at least one vocab
        code, o, e = run("export", self.m, "--out", wl, "--reviewer1-id", "a", "--reviewer2-id", "b")
        self.assertEqual(code, 1)
        self.assertIn("outside the Malayalam block", e)
        code, o, e = run("export", self.m, "--out", wl, "--reviewer1-id", "a", "--reviewer2-id", "b", "--allow-non-malayalam")
        self.assertEqual(code, 0, e)

    def test_not_a_master(self):
        self.propose()
        out = os.path.join(self.d, "copies")
        run("split", self.m, "--out-dir", out)
        c1 = os.path.join(out, "review_copy_reviewer1.xlsx")
        self.assertNotEqual(run("split", c1, "--out-dir", os.path.join(self.d, "z"))[0], 0)
        self.assertNotEqual(run("status", c1)[0], 0)


@unittest.skipUnless(os.path.exists(os.path.join(BUILD, "s2w_make_fixture")) and os.path.exists(os.path.join(BUILD, "s2w_dataset_check")),
                     "C tools not built")
class ValidatorRoundTrip(unittest.TestCase):
    """Export from a (synthetic) review sheet and feed the result to the C dataset validator."""

    def test_export_is_accepted_by_dataset_check(self):
        d = tempfile.mkdtemp()
        try:
            fx = os.path.join(d, "fx")
            subprocess.run([os.path.join(BUILD, "s2w_make_fixture"), fx], check=True, stdout=subprocess.DEVNULL)
            with open(os.path.join(fx, "wordlist.csv"), encoding="utf-8") as f:
                f.readline()
                words = list(csv.DictReader(f))
            m = os.path.join(d, "m.xlsx")
            self.assertEqual(run("template", m, "--vocab-slots", "30", "--oov-slots", "10")[0], 0)
            wb = load_workbook(m)
            ws = wb["Master"]
            hdr = {c.value: c.column for c in ws[1]}
            col = lambda k: hdr[rs.HEAD[k]]  # noqa: E731
            for i, w in enumerate(words, start=2):
                vals = dict(id=w["word_id"], role=w["role"], word=w["label"], meaning=w["meaning_en"], roman=w["romanisation"],
                            pair=w["pair_id"] or None, ctype=w["contrast_type"], r1_dec="Accept", r2_dec="Accept", r1_fam="Yes", r2_fam="Yes",
                            r1_con="Yes" if w["pair_id"] else None, r2_con="Yes" if w["pair_id"] else None, final=w["label"], approved="Yes")
                for k, v in vals.items():
                    ws.cell(row=i, column=col(k), value=v)
            wb.save(m)
            wl = os.path.join(d, "wordlist.csv")
            code, o, e = run("export", m, "--out", wl, "--reviewer1-id", "rev-a", "--reviewer2-id", "rev-b", "--allow-non-malayalam")
            self.assertEqual(code, 0, e)
            shutil.copy(wl, os.path.join(fx, "wordlist.csv"))
            r = subprocess.run([os.path.join(BUILD, "s2w_dataset_check"), fx, "--no-audio"], capture_output=True, text=True)
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
            self.assertIn("result=VALID", r.stdout)
        finally:
            shutil.rmtree(d, ignore_errors=True)


if __name__ == "__main__":
    unittest.main(verbosity=1)
