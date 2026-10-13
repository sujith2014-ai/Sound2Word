"""Tests for tools/s2w_viewer.py (stdlib unittest). Uses the SYNTHETIC fixture only."""
import hashlib
import json
import shutil
import signal
import socket
import time
import math
import os
import struct
import subprocess
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
import wave

BUILD = os.environ.get("S2W_BUILD", "build")
NEEDED = ("s2w_make_fixture", "s2w_eval", "s2w_features", "s2w_qa", "s2w_match")
if not all(os.path.exists(os.path.join(BUILD, n)) for n in NEEDED):
    print("viewer tests skipped: C tools not built")
    sys.exit(0)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import s2w_viewer as v  # noqa: E402


def sha(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def tree_hash(root):
    h = hashlib.sha256()
    for d, _, fs in sorted(os.walk(root)):
        for n in sorted(fs):
            h.update(n.encode())
            h.update(sha(os.path.join(d, n)).encode())
    return h.hexdigest()


class Pure(unittest.TestCase):
    def test_peaks_match_brute_force(self):
        s = [(-1) ** i * (i * 37 % 30000) for i in range(1000)]
        p = v.peaks(s, 10)
        self.assertEqual(len(p), 10)
        for b, (lo, hi) in enumerate(p):
            seg = s[b * 100:(b + 1) * 100]
            self.assertAlmostEqual(lo, min(seg) / 32768.0)
            self.assertAlmostEqual(hi, max(seg) / 32768.0)

    def test_peaks_short_input(self):
        self.assertEqual(len(v.peaks([5, -5, 7], 1200)), 3)

    def test_f0_of_150hz_tone(self):
        if v.np is None:
            self.skipTest("numpy not installed")
        s = [int(8000 * math.sin(2 * math.pi * 150 * i / 16000)) for i in range(16000)]
        f = [x for x in v.f0_track(s, 16000) if x]
        self.assertTrue(len(f) > 50)
        self.assertTrue(all(abs(x - 150) < 6 for x in f))

    def test_segments_of_three_tone_signal(self):
        # mel-like matrix: strongest bin 5, then 15, then 25, with silence around
        mat = [[-20.0] * 40 for _ in range(10)]
        for b, n in ((5, 20), (15, 20), (25, 20)):
            for _ in range(n):
                r = [-20.0] * 40
                r[b] = 5.0
                mat.append(r)
        mat += [[-20.0] * 40 for _ in range(10)]
        seg = v.estimate_segments(mat)
        self.assertEqual(len(seg), 3)
        self.assertAlmostEqual(seg[0]["start"], 0.1075)               # frame 10: centre 0.1125, owns +/-5 ms
        for a, b in zip(seg, seg[1:]):
            self.assertAlmostEqual(a["end"], b["start"])                # adjacent segments abut


class Server(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.ds = os.path.join(cls.tmp, "ds")
        subprocess.run([os.path.join(BUILD, "s2w_make_fixture"), cls.ds], check=True, capture_output=True)
        subprocess.run([os.path.join(BUILD, "s2w_dataset_check"), cls.ds, "--emit-hashes"], capture_output=True)
        cfg = os.path.join(cls.tmp, "base.cfg")
        with open(cfg, "w") as f:
            f.write("#s2w-settings v1\n")
        subprocess.run([os.path.join(BUILD, "s2w_eval"), "run", "--dataset", cls.ds, "--settings", cfg, "--split", "calib", "--out-dir", os.path.join(cls.tmp, "runs"),
                        "--run-id", "t"], check=True, capture_output=True)
        cls.run_dir = os.path.join(cls.tmp, "runs", "t")
        cls.app = v.App(cls.ds, cls.run_dir, BUILD)
        cls.srv = v.make_server(cls.app, 0)
        cls.port = cls.srv.server_address[1]
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()
        cls.before = (tree_hash(cls.ds), tree_hash(cls.run_dir))

    @classmethod
    def tearDownClass(cls):
        cls.srv.shutdown()
        cls.srv.server_close()

    def get(self, path):
        with urllib.request.urlopen("http://127.0.0.1:%d%s" % (self.port, path)) as r:
            return r.read()

    def err(self, path, method="GET"):
        try:
            urllib.request.urlopen(urllib.request.Request("http://127.0.0.1:%d%s" % (self.port, path), method=method, data=b"x" if method != "GET" else None))
        except urllib.error.HTTPError as e:
            return e.code
        return 200

    def test_index_counts_and_synthetic_banner_data(self):
        idx = json.loads(self.get("/api/index"))
        self.assertEqual(idx["kind"], "synthetic")
        self.assertEqual(len(idx["queries"]), len(self.app.queries))

    def test_recording_is_real_audio_analysis(self):
        rid = self.app.queries[0]["recording_id"]
        d = json.loads(self.get("/api/recording/" + rid))
        with wave.open(self.app.wav_path(rid)) as w:
            self.assertEqual(d["n_samples"], w.getnframes())
            self.assertAlmostEqual(d["duration"], w.getnframes() / w.getframerate())
        self.assertEqual(d["spec"]["cols"], 40)
        self.assertEqual(len(d["segments"]), 3)               # SYNTH_V1 has three tones
        self.assertIsNone(d["f0"])                            # F0 hidden for synthetic tones
        self.assertEqual(d["qa"]["verdict"], "ok")

    def test_audio_endpoint_returns_identical_bytes(self):
        rid = self.app.queries[0]["recording_id"]
        with open(self.app.wav_path(rid), "rb") as f:
            self.assertEqual(self.get("/audio/" + rid), f.read())

    def test_range_requests_for_seekable_playback(self):
        rid = self.app.queries[0]["recording_id"]
        with open(self.app.wav_path(rid), "rb") as f:
            full = f.read()

        def rq(rng):
            req = urllib.request.Request("http://127.0.0.1:%d/audio/%s" % (self.port, rid), headers={"Range": rng} if rng else {})
            try:
                with urllib.request.urlopen(req) as r:
                    return r.status, r.read(), dict(r.headers)
            except urllib.error.HTTPError as e:
                return e.code, b"", dict(e.headers)
        st, body, h = rq(None)
        self.assertEqual((st, h.get("Accept-Ranges")), (200, "bytes"))
        st, body, h = rq("bytes=100-199")
        self.assertEqual((st, body), (206, full[100:200]))
        self.assertEqual(h["Content-Range"], "bytes 100-199/%d" % len(full))
        self.assertEqual(rq("bytes=%d-" % (len(full) - 10))[1], full[-10:])
        self.assertEqual(rq("bytes=-20")[1], full[-20:])
        for bad in ("bytes=%d-" % (len(full) + 5), "bytes=9-3", "items=0-1", "bytes=0-1,5-6", "bytes=x-y"):
            self.assertEqual(rq(bad)[0], 416, bad)

    def test_unknown_ids_and_traversal_refused(self):
        for p in ("/audio/nope", "/audio/..%2F..%2Fetc%2Fpasswd", "/api/recording/..%2Frecordings.csv", "/api/match/nope"):
            self.assertEqual(self.err(p), 404, p)

    def test_match_ranks_same_word_first_and_agrees_with_engine(self):
        q = next(x for x in self.app.queries if x["word_id"] == "SYNTH_V1")
        m = json.loads(self.get("/api/match/" + q["recording_id"]))
        self.assertEqual(m["ranks"][0]["word"], "SYNTH_V1")
        self.assertEqual(self.app.label_to_word[q["top1"]], m["ranks"][0]["word"])
        self.assertTrue(all(m["ranks"][i]["dist"] <= m["ranks"][i + 1]["dist"] for i in range(len(m["ranks"]) - 1)))

    def test_read_only(self):
        for method in ("POST", "PUT", "DELETE"):
            self.assertEqual(self.err("/api/index", method), 405)
        self.get("/api/recording/" + self.app.queries[0]["recording_id"])
        self.get("/api/match/" + self.app.queries[0]["recording_id"])
        self.assertEqual((tree_hash(self.ds), tree_hash(self.run_dir)), self.before)

    def test_verification_cannot_approve(self):
        v_ = json.loads(self.get("/api/verification"))
        self.assertFalse(v_["available"])                     # no workbook supplied
        self.assertIn("cannot approve", v_["note"])


def free_port():
    with socket.socket() as sk:
        sk.bind(("127.0.0.1", 0))
        return sk.getsockname()[1]


def fresh_fixture():
    tmp = tempfile.mkdtemp()
    ds = os.path.join(tmp, "ds")
    subprocess.run([os.path.join(BUILD, "s2w_make_fixture"), ds], check=True, capture_output=True)
    subprocess.run([os.path.join(BUILD, "s2w_dataset_check"), ds, "--emit-hashes"], capture_output=True)
    cfg = os.path.join(tmp, "base.cfg")
    with open(cfg, "w") as f:
        f.write("#s2w-settings v1\n")
    subprocess.run([os.path.join(BUILD, "s2w_eval"), "run", "--dataset", ds, "--settings", cfg, "--split", "calib", "--out-dir", os.path.join(tmp, "runs"),
                    "--run-id", "t"], check=True, capture_output=True)
    return tmp, ds, os.path.join(tmp, "runs", "t")


VIEWER = os.path.join(os.path.dirname(os.path.abspath(__file__)), "s2w_viewer.py")


class Cli(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp, cls.ds, cls.run_dir = fresh_fixture()

    def viewer(self, *a):
        return subprocess.run([sys.executable, VIEWER, "--build", BUILD] + list(a), capture_output=True, text=True, timeout=30)

    def test_invalid_paths_give_clean_error_not_traceback(self):
        for args in (["--dataset", "/no/such", "--run", self.run_dir], ["--dataset", self.ds, "--run", "/no/such"]):
            r = self.viewer(*args, "--port", str(free_port()))
            self.assertEqual(r.returncode, 2)
            self.assertIn("error:", r.stderr)
            self.assertNotIn("Traceback", r.stderr)

    def test_malformed_queries_file_named_in_error(self):
        bad = os.path.join(self.tmp, "badrun")
        shutil.copytree(self.run_dir, bad)
        with open(os.path.join(bad, "queries.csv"), "w") as f:
            f.write("foo,bar\n1,2\n")
        r = self.viewer("--dataset", self.ds, "--run", bad, "--port", str(free_port()))
        self.assertEqual(r.returncode, 2)
        self.assertIn("missing required column", r.stderr)
        self.assertNotIn("Traceback", r.stderr)

    def test_port_in_use_is_reported(self):
        with socket.socket() as sk:
            sk.bind(("127.0.0.1", 0))
            sk.listen()
            r = self.viewer("--dataset", self.ds, "--run", self.run_dir, "--port", str(sk.getsockname()[1]))
        self.assertEqual(r.returncode, 2)
        self.assertIn("cannot listen", r.stderr)
        self.assertNotIn("Traceback", r.stderr)

    def test_serves_page_then_sigterm_exits_cleanly_and_frees_port(self):
        port = free_port()
        proc = subprocess.Popen([sys.executable, VIEWER, "--build", BUILD, "--dataset", self.ds, "--run", self.run_dir, "--port", str(port)],
                                stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
        try:
            page = None
            for _ in range(50):
                try:
                    page = urllib.request.urlopen("http://127.0.0.1:%d/" % port).read().decode()
                    break
                except OSError:
                    time.sleep(0.1)
            self.assertIsNotNone(page)
            self.assertIn("Sound2Word", page)
            proc.send_signal(signal.SIGTERM)
            self.assertEqual(proc.wait(timeout=10), 0)
        finally:
            if proc.poll() is None:
                proc.kill()
                proc.wait()
            proc.stdout.close()
            proc.stderr.close()
        with socket.socket() as sk:                       # nothing is listening any more
            sk.settimeout(2)
            self.assertNotEqual(sk.connect_ex(("127.0.0.1", port)), 0)


class BadInputs(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp, cls.ds, cls.run_dir = fresh_fixture()
        cls.app = v.App(cls.ds, cls.run_dir, BUILD)
        cls.srv = v.make_server(cls.app, 0)
        cls.port = cls.srv.server_address[1]
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()
        cls.rid = cls.app.queries[0]["recording_id"]
        cls.wav = cls.app.wav_path(cls.rid)
        with open(cls.wav, "rb") as f:
            cls.good = f.read()

    @classmethod
    def tearDownClass(cls):
        cls.srv.shutdown()
        cls.srv.server_close()

    def tearDown(self):
        with open(self.wav, "wb") as f:
            f.write(self.good)
        self.app.build = BUILD
        self.app.review = None

    def get(self, path):
        try:
            with urllib.request.urlopen("http://127.0.0.1:%d%s" % (self.port, path)) as r:
                return r.status, r.read().decode()
        except urllib.error.HTTPError as e:
            return e.code, e.read().decode()

    def assert_clean(self, body):
        self.assertNotIn(self.tmp, body)                  # no local paths leak to the page
        self.assertNotIn("Traceback", body)

    def test_missing_audio(self):
        os.unlink(self.wav)
        for p in ("/api/recording/", "/audio/", "/api/match/"):
            code, body = self.get(p + self.rid)
            self.assertEqual(code, 404, p)
            self.assert_clean(body)

    def test_corrupt_wav(self):
        with open(self.wav, "wb") as f:
            f.write(b"garbage" * 50)
        code, body = self.get("/api/recording/" + self.rid)
        self.assertEqual(code, 422)
        self.assert_clean(body)

    def test_stereo_wav(self):
        with wave.open(self.wav, "wb") as w:
            w.setnchannels(2)
            w.setsampwidth(2)
            w.setframerate(16000)
            w.writeframes(b"\0\0" * 3200)
        code, body = self.get("/api/recording/" + self.rid)
        self.assertEqual(code, 422)
        self.assertIn("16-bit mono", body)

    def test_wav_with_no_samples(self):
        with open(self.wav, "wb") as f:
            f.write(self.good[:44])
        code, body = self.get("/api/recording/" + self.rid)
        self.assertEqual(code, 422)
        self.assert_clean(body)

    def test_missing_analysis_tools(self):
        self.app.build = os.path.join(self.tmp, "no-build-dir")
        code, body = self.get("/api/recording/" + self.rid)
        self.assertEqual(code, 503)
        self.assertIn("make all", body)
        self.assert_clean(body)

    def test_malformed_review_workbook_shows_nothing_as_verified(self):
        bad = os.path.join(self.tmp, "bad.xlsx")
        with open(bad, "wb") as f:
            f.write(b"not a workbook")
        self.app.review = bad
        code, body = self.get("/api/verification")
        self.assertEqual(code, 200)
        d = json.loads(body)
        self.assertFalse(d["available"])
        self.assertEqual(d["words"], [])
        self.assert_clean(body)

    def test_missing_review_workbook(self):
        self.app.review = os.path.join(self.tmp, "absent.xlsx")
        self.assertFalse(json.loads(self.get("/api/verification")[1])["available"])

    def test_still_serves_after_all_failures(self):
        self.assertEqual(self.get("/api/recording/" + self.rid)[0], 200)


class ReviewWorkbookUntouched(unittest.TestCase):
    def test_viewing_workbook_changes_nothing_and_approves_nothing(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl not installed")
        tmp, ds, run_dir = fresh_fixture()
        master = os.path.join(tmp, "master.xlsx")
        subprocess.run([sys.executable, os.path.join(os.path.dirname(VIEWER), "review_sheet.py"), "template", master, "--vocab-slots", "2", "--oov-slots", "0"],
                       check=True, capture_output=True)
        import openpyxl
        wb = openpyxl.load_workbook(master)
        ws = wb["Master"]
        ws.cell(2, 3).value = "X"
        ws.cell(2, 5).value = "meaning"
        wb.save(master)
        before = sha(master)
        app = v.App(ds, run_dir, BUILD, master)
        out = app.verification()
        self.assertTrue(out["available"])
        self.assertEqual(sha(master), before)
        text = json.dumps(out).lower()
        self.assertNotIn("approved", text.replace("cannot approve", ""))
        self.assertTrue(all("verified" not in w["status"].lower() or "pending" in w["status"].lower() for w in out["words"]))


if __name__ == "__main__":
    unittest.main()
